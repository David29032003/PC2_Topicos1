import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

from comun import exportar, nueva_presentacion

aqui = Path(__file__).resolve().parent
res = aqui.parent / "results"

filas = [json.loads(l) for l in open(res / "exp1_consulta.jsonl", encoding="utf-8")]
q0 = {f["query_id"]: f for f in filas if f["condicion"] == "Q0"}
q1 = {f["query_id"]: f for f in filas if f["condicion"] == "Q1"}
ids = sorted(q0)


def num(x):
    if x == 0:
        return "0"
    if x == 1:
        return "1"
    return f"{x:.2f}".replace(".", ",")


def media(d, k):
    return sum(f[k] for f in d.values()) / len(d)


prs = None  # lo fija agregar()


def slide_titulo(titulo):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.7), Inches(0.4), Inches(11.9), Inches(1.0)
    t.text_frame.text = titulo
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(34)
    t.text_frame.paragraphs[0].alignment = 1
    return s


def bullets(s, items, x=0.7, y=1.6, w=11.9, h=5.2, size=22):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        nivel = 1 if it.startswith("  ") else 0
        p.text = ("- " if nivel else "• ") + it.strip()
        p.level = nivel
        p.space_after = Pt(10)
        for r in p.runs:
            r.font.size = Pt(size - 3 * nivel)
    return tb


def caja(s, x, y, w, h, texto, size=16):
    b = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.background()
    b.line.color.rgb = __import__("pptx").dml.color.RGBColor(0, 0, 0)
    b.line.width = Pt(1.25)
    b.shadow.inherit = False
    tf = b.text_frame
    tf.word_wrap = True
    for i, linea in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = linea
        p.alignment = 2
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.color.rgb = __import__("pptx").dml.color.RGBColor(0, 0, 0)
    return b


def flecha(s, x1, y1, x2, y2):
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = __import__("pptx").dml.color.RGBColor(0, 0, 0)
    c.line.width = Pt(1.5)
    from lxml import etree
    ln = c.line._get_or_add_ln()
    e = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    e.set("type", "triangle")


def tabla(s, datos, x, y, w, anchos, alto=0.42, size=14):
    gf = s.shapes.add_table(len(datos), len(datos[0]), Inches(x), Inches(y), Inches(w), Inches(alto * len(datos)))
    t = gf.table
    for j, a in enumerate(anchos):
        t.columns[j].width = Inches(a)
    for i, fila in enumerate(datos):
        t.rows[i].height = Inches(alto)
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            c.text = str(v)
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
    return t


def notas(s, texto):
    s.notes_slide.notes_text_frame.text = texto


def agregar(p, portada=True):
    global prs
    prs = p
    if portada:
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = "Experimento 1: consulta original vs consulta estructurada"
        s.placeholders[1].text = "PC2 - Tópicos de Ciencia de la Computación 1"
        s.shapes.title.left, s.shapes.title.top, s.shapes.title.width, s.shapes.title.height = Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.8)
        s.placeholders[1].left, s.placeholders[1].top, s.placeholders[1].width, s.placeholders[1].height = Inches(0.9), Inches(4.1), Inches(11.5), Inches(1.0)


    # 2. de que trata
    s = slide_titulo("¿De qué trata el experimento?")
    bullets(s, [
        "El auditor escribe una necesidad de información en lenguaje natural",
        "  Ejemplo: ¿Qué empresa ganó la compra de reactivos para el equipo Mindray?",
        "Antes de buscar, un LLM la convierte en una consulta estructurada (un JSON)",
        "Pregunta: ¿esa transformación conserva, mejora o empeora lo que se quería encontrar?",
        "No se exige demostrar que la consulta del LLM sea mejor, solo observar qué pasó",
    ])
    notas(s, "Es la parte de la Semana 3 (salida estructurada y validación) conectada con la Semana 4 (recuperación).")

    # 3. que se compara
    s = slide_titulo("Qué se compara")
    caja(s, 4.7, 1.5, 3.9, 0.9, "Necesidad del auditor\n(una de las 5)")
    flecha(s, 5.6, 2.4, 3.2, 3.1)
    flecha(s, 7.7, 2.4, 10.1, 3.1)
    caja(s, 1.0, 3.1, 4.4, 1.0, "Q0\nLa necesidad tal como se escribió")
    caja(s, 7.9, 3.1, 4.4, 1.0, "Q1\nCampo query del JSON del LLM")
    flecha(s, 10.1, 4.1, 10.1, 4.6)
    caja(s, 7.9, 4.6, 4.4, 0.9, "Qwen2.5-0.5B + validación\n(parseable, schema, intención)")
    flecha(s, 3.2, 4.1, 5.2, 5.9)
    flecha(s, 10.1, 5.5, 8.2, 5.9)
    caja(s, 4.7, 5.9, 3.9, 0.9, "Mismo recuperador\n(e5 + FAISS, chunks de 180 palabras)")
    notas(s, "Las dos rutas terminan en el mismo recuperador, así que cualquier diferencia solo puede venir de la consulta.")

    # 4. fijo y metricas
    s = slide_titulo("Qué se mantiene igual y qué se mide")
    bullets(s, [
        "Se mantiene igual: corpus (10 fichas), chunks de 180 palabras, modelo e5-small, FAISS, top-k y qrels",
        "Solo cambia la consulta, por eso la diferencia se puede atribuir a ella",
        "Sobre el JSON del LLM se mide:",
        "  schema_valid_rate: ¿cumple el formato?",
        "  intención correcta: ¿la intención coincide con la esperada?",
        "Sobre la recuperación se mide:",
        "  Recall@1, Recall@3 y Recall@5 de Q0 y de Q1",
        "  Jaccard del top-3: cuánto coinciden los fragmentos recuperados por ambas",
    ])

    # 5. resultados: formato del json
    n_ok = sum(q1[i]["llm"]["intent_correcto"] for i in ids)
    n_sv = sum(q1[i]["llm"]["schema_valid"] for i in ids)
    n_ex = sum(q1[i]["llm"]["json_exacto"] for i in ids)
    n_pa = sum(q1[i]["llm"]["json_recuperado"] for i in ids)
    n_ab = sum(bool(q1[i]["llm"]["abstain"]) for i in ids)
    n = len(ids)

    s = slide_titulo("Resultados: el formato del JSON")
    datos = [["Métrica", "Valor", "Qué mide"],
             ["JSON exacto", f"{n_ex}/{n}", "el texto completo es un JSON"],
             ["JSON parseable", f"{n_pa}/{n}", "se puede extraer el JSON del texto"],
             ["schema_valid_rate", f"{n_sv}/{n}", "cumple claves, tipos y valores permitidos"],
             ["Abstención", f"{n_ab}/{n}", "el modelo considera que la necesidad sí se puede responder con estas fichas, así que genera la consulta"]]
    tabla(s, datos, 0.7, 1.6, 11.9, [3.2, 1.7, 7.0], size=16)
    tb = bullets(s, [
        "Interpretación",
        "  El modelo cumple el formato: las 5 salidas son válidas",
        "  El JSON exacto es 0 porque siempre lo envolvió en un bloque de código; al parsear se recupera",
        "  En formato no hay problema, pero el schema solo revisa la forma, no si el contenido es correcto",
    ], y=4.2, h=2.8, size=22)
    tb.text_frame.paragraphs[0].runs[0].font.bold = True

    # 6. resultados: contenido (intencion)
    s = slide_titulo("Resultados: la intención del JSON")
    esp = {f["query_id"]: f["intent_gold"] for f in q1.values()}
    datos = [["Necesidad", "Intención esperada", "Intención del LLM", "¿Correcta?"]]
    for i in ids:
        l = q1[i]["llm"]
        datos.append([i, esp[i], l["intent"], "sí" if l["intent_correcto"] else "no"])
    tabla(s, datos, 0.7, 1.6, 11.9, [1.6, 4.1, 4.1, 2.1], size=15)
    tb = bullets(s, [
        "Interpretación",
        f"  Intención correcta en {n_ok} de {n}: formato válido no garantiza contenido correcto",
        "  En N5 (¿quién ganó?) el modelo respondió marca_con_equivalente (posible causa: la palabra Mindray, una marca)",
        "  Un sistema que filtre por intención buscaría en el lugar equivocado en N3 y N5",
    ], y=4.4, h=2.8, size=20)
    tb.text_frame.paragraphs[0].runs[0].font.bold = True

    # 7. resultados: recuperacion
    s = slide_titulo("Resultados: la recuperación")
    cd = CategoryChartData()
    cd.categories = ids
    cd.add_series("Q0", [q0[i]["recall@3"] for i in ids])
    cd.add_series("Q1", [q1[i]["recall@3"] for i in ids])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(1.5), Inches(7.0), Inches(4.8), cd)
    ch = gf.chart
    ch.has_title = True
    ch.chart_title.text_frame.text = "Recall@3 por necesidad"
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    pl = ch.plots[0]
    pl.gap_width = 40
    for serie, d in zip(pl.series, (q0, q1)):
        for punto, i in zip(serie.points, ids):
            tf = punto.data_label.text_frame
            tf.text = num(d[i]["recall@3"])
            tf.paragraphs[0].runs[0].font.size = Pt(11)
    ch.value_axis.maximum_scale = 1.1
    ch.value_axis.minimum_scale = 0
    ch.value_axis.has_major_gridlines = False
    datos = [["Consulta", "R@1", "R@3", "R@5"]]
    for nombre, d in (("Q0", q0), ("Q1", q1)):
        datos.append([nombre] + [num(media(d, f"recall@{k}")) for k in (1, 3, 5)])
    tabla(s, datos, 8.2, 1.7, 4.4, [1.3, 1.0, 1.0, 1.1], size=16)
    jac = media(q1, "jaccard_top3_con_Q0")
    bullets(s, [
        "Q0 y Q1 dan los mismos valores",
        f"Jaccard medio del top-3 = {num(jac)}",
        "Solo N2 cambia de fragmentos",
    ], x=8.2, y=3.5, w=4.5, h=3.2, size=18)

    # 8. interpretacion de la recuperacion
    s = slide_titulo("Interpretación: la recuperación")
    tb = bullets(s, [
        "Q0 y Q1 recuperan lo mismo",
        "  El LLM copió la necesidad en 4 de 5 casos, por eso el Recall no cambia",
        "  Jaccard 0,90: casi todos los top-3 son iguales (en N2 comparten 2 de 4 fragmentos)",
        "El Recall es bajo en las dos consultas",
        f"  Con 5 resultados queda sin encontrar casi la mitad de la evidencia (R@5 = {num(media(q0, 'recall@5'))})",
        "  El problema está en el recuperador y no en la transformación de la consulta",
        "Recall@1 ≤ Recall@3 ≤ Recall@5, como debe ser: más resultados cubren más evidencia",
    ], size=22)

    # 9. por necesidad
    s = slide_titulo("Interpretación: por necesidad")
    datos = [["Necesidad", "R@3 (Q0 y Q1)", "Por qué"],
             ["N1", num(q0["N1"]["recall@3"]), "encuentra la ficha Cisco, no la de Agilent"],
             ["N2", num(q0["N2"]["recall@3"]), "la evidencia es la frase corta \"o equivalente\", que el modelo no destaca"],
             ["N3", num(q0["N3"]["recall@3"]), "encuentra la ficha Mindray, no la de Epson"],
             ["N4", num(q0["N4"]["recall@3"]), "trae un proceso con un solo postor de los tres"],
             ["N5", num(q0["N5"]["recall@3"]), "pregunta por un caso concreto con palabras que están en la ficha"]]
    tabla(s, datos, 0.7, 1.6, 11.9, [1.7, 2.6, 7.6], size=16, alto=0.5)
    bullets(s, [
        "Las necesidades con evidencia en varias fichas (N1, N3, N4) quedan a medias",
        "Las que dependen de una frase corta (N2) quedan en 0",
    ], y=5.0, h=2.0, size=20)

    # 10. que paso con la consulta
    s = slide_titulo("Qué pasó con la consulta")
    bullets(s, [
        "En 4 de 5 necesidades el LLM copió la necesidad tal cual como query",
        "En N2 la reformuló:",
        "  Original: Quiero encontrar los requerimientos que mencionan una marca pero permiten ofertar un producto equivalente",
        "  Q1: ¿Qué requerimientos mencionan una marca pero permiten ofertar un producto equivalente?",
        "  Cambió uno de los tres fragmentos del top-3 (Jaccard 0,5), pero ninguno era relevante y el R@3 siguió en 0",
        "En N5 el JSON fue válido pero la intención salió marca_con_equivalente, siendo una pregunta por el ganador",
    ], size=20)
    notas(s, "El schema solo verifica que el valor esté dentro de los permitidos, no que sea el correcto.")

    # 11. conclusion
    s = slide_titulo("Conclusión")
    bullets(s, [
        "La consulta estructurada conserva la necesidad, pero no la mejora",
        "  El modelo pequeño no la reescribe con el vocabulario de las fichas (\"buena pro\", \"número de postores\", \"o equivalente\")",
        "Un JSON válido no garantiza una intención correcta",
        f"  schema_valid_rate = 1 mide el formato; la intención fue correcta en {n_ok} de {n}",
        "El Recall bajo viene del recuperador, no de la transformación",
        "  N2 queda en 0 con Q0 y con Q1",
        "Con 5 necesidades solo se pueden sacar conclusiones sobre estos casos",
    ], size=22)


def construir():
    p = nueva_presentacion()
    agregar(p)
    salida = aqui / "Experimento1.pptx"
    p.save(salida)
    exportar(salida, aqui / "Experimento1.pdf", aqui / "_png_exp1")
    print("listo", salida)


if __name__ == "__main__":
    construir()
