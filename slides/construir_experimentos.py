from pathlib import Path

from pptx.util import Inches, Pt

import construir_exp1 as e1
import construir_exp2 as e2
from comun import exportar, nueva_presentacion

aqui = Path(__file__).resolve().parent
diagrama = aqui / "diagrama_flujo_pc2.png"

prs = nueva_presentacion()

# portada
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "PC2 - Tópicos de Ciencia de la Computación 1"
s.placeholders[1].text = "Experimento 1: consulta estructurada  |  Experimento 2: tamaño del chunk"
s.shapes.title.left, s.shapes.title.top, s.shapes.title.width, s.shapes.title.height = Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.8)
s.placeholders[1].left, s.placeholders[1].top, s.placeholders[1].width, s.placeholders[1].height = Inches(0.9), Inches(4.1), Inches(11.5), Inches(1.0)

# diagrama de flujo
e1.prs = prs  # los helpers de e1 trabajan sobre esta presentacion
s = e1.slide_titulo("Diagrama de flujo de la PC2")
alto = 6.0
ancho = alto * 839 / 870
s.shapes.add_picture(str(diagrama), Inches((13.333 - ancho) / 2), Inches(1.35), Inches(ancho), Inches(alto))
e1.notas(s, "Dos ramas que se juntan en FAISS. A la izquierda los documentos: se fragmentan, e5 convierte cada "
            "fragmento en un vector y FAISS los guarda. Arriba la consulta: el usuario escribe la necesidad, el LLM "
            "la convierte en JSON, se valida, e5 la convierte en vector y FAISS busca los mas parecidos. Abajo se "
            "evalua: top-k contra los qrels da el Recall@k.")

# como leer el diagrama
s = e1.slide_titulo("Cómo leer el diagrama")
e1.bullets(s, [
    "Rama de los documentos (izquierda)",
    "  Fichas → se fragmentan en chunks → e5 convierte cada chunk en un vector → FAISS los guarda",
    "Rama de la consulta (arriba)",
    "  Necesidad → prompt (system + user) → LLM → JSON (query, intent, abstain) → validación → e5 → vector",
    "Búsqueda y evaluación (abajo)",
    "  FAISS compara la consulta con todos los chunks → scores e ids → top-k → Recall@k con los qrels",
    "Qué cambia cada experimento",
    "  Experimento 1: la rama de arriba (consulta original o consulta del LLM)",
    "  Experimento 2: la rama de la izquierda (tamaño de los chunks: 100, 180 o 320 palabras)",
], size=20)
e1.notas(s, "El JSON y los qrels que aparecen en el diagrama son un ejemplo generico. En la PC2 las intenciones son "
            "como marca_sin_equivalente o competencia, y los qrels apuntan a passages (por ejemplo D05-P06), no a "
            "documentos completos.")

# los dos experimentos completos; la portada de cada uno hace de separador
e1.agregar(prs, portada=True)
e2.agregar(prs, portada=True)

salida = aqui / "Experimentos_PC2.pptx"
prs.save(salida)
exportar(salida, aqui / "Experimentos_PC2.pdf", aqui / "_png_exps")
print("listo", salida, "| diapositivas:", len(prs.slides))
