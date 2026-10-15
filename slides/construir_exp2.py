import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

from comun import exportar, nueva_presentacion

aqui = Path(__file__).resolve().parent
res = aqui.parent / "results"

filas = [json.loads(l) for l in open(res / "exp2_chunking.jsonl", encoding="utf-8")]
resumen = json.load(open(res / "resumen.json", encoding="utf-8"))
conds = ["small", "baseline", "large"]
palabras = {"small": 100, "baseline": 180, "large": 320}
ids = sorted({f["query_id"] for f in filas})
r = {(f["condicion"], f["query_id"]): f for f in filas}
chunks = {c: [json.loads(l) for l in open(res / f"chunks_{c}.jsonl", encoding="utf-8")] for c in conds}


def num(x):
    if x == 0:
        return "0"
    if x == 1:
        return "1"
    return f"{x:.2f}".replace(".", ",")


def media(cond, k):
    return sum(r[(cond, i)][f"recall@{k}"] for i in ids) / len(ids)


prs = None  # lo fija agregar()


def slide_titulo(titulo):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.7), Inches(0.4), Inches(11.9), Inches(1.0)
    t.text_frame.text = titulo
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(34)
    t.text_frame.paragraphs[0].alignment = 1
    return s


def bullets(s, items, x=0.7, y=1.6, w=11.9, h=5.2, size=22):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        nivel = 1 if it.startswith("  ") else 0
        p.text = ("- " if nivel else "• ") + it.strip()
        p.level = nivel
        p.space_after = Pt(10)
        for run in p.runs:
            run.font.size = Pt(size - 3 * nivel)
    return tb


def tabla(s, datos, x, y, w, anchos, alto=0.42, size=14):
    gf = s.shapes.add_table(len(datos), len(datos[0]), Inches(x), Inches(y), Inches(w), Inches(alto * len(datos)))
    t = gf.table
    for j, a in enumerate(anchos):
        t.columns[j].width = Inches(a)
    for i, fila in enumerate(datos):
        t.rows[i].height = Inches(alto)
        for j, v in enumerate(fila):
            c = t.cell(i, j)
            c.text = str(v)
            for p in c.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(size)
    return t


def notas(s, texto):
    s.notes_slide.notes_text_frame.text = texto


def pasajes(c):
    p = c["passage_ids"]
    return f"{p[0][-3:]} a {p[-1][-3:]}" if len(p) > 1 else p[0][-3:]


def agregar(p, portada=True):
    global prs
    prs = p
    if portada:
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = "Experimento 2: tamaño del chunk"
        s.placeholders[1].text = "PC2 - Tópicos de Ciencia de la Computación 1"
        s.shapes.title.left, s.shapes.title.top, s.shapes.title.width, s.shapes.title.height = Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.8)
        s.placeholders[1].left, s.placeholders[1].top, s.placeholders[1].width, s.placeholders[1].height = Inches(0.9), Inches(4.1), Inches(11.5), Inches(1.0)


    # 2. de que trata
    s = slide_titulo("¿De qué trata el experimento?")
    bullets(s, [
        "Pregunta: ¿cómo cambia la recuperación si solo cambia el tamaño de los chunks?",
        "No interviene el LLM ni ningún prompt: es un experimento sobre el recuperador",
        "Una ficha se divide en passages (piezas pequeñas) y un chunk agrupa passages consecutivos",
        "  Chunk pequeño: más chunks, cada uno con menos contexto",
        "  Chunk grande: menos chunks, cada uno con casi toda la ficha",
        "Hipótesis: un chunk pequeño aísla la evidencia, pero puede separarla de lo que la identifica",
    ])
    notas(s, "Es la parte de la Semana 4: segmentación, embeddings, FAISS y Recall@k.")

    # 3. que se compara y que se mantiene
    s = slide_titulo("Qué se compara y qué se mantiene igual")
    datos = [["Condición", "Tamaño objetivo", "Chunks resultantes"]]
    for c in conds:
        datos.append([c, f"{palabras[c]} palabras", str(len(chunks[c]))])
    tabla(s, datos, 0.7, 1.6, 7.0, [2.2, 2.6, 2.2], size=18, alto=0.5)
    bullets(s, [
        "Se mantiene igual:",
        "  Las mismas fichas, las mismas 5 consultas (Q0) y los mismos qrels",
        "  Solapamiento de 1 passage entre chunks vecinos",
        "  Modelo e5-small, normalización, FAISS y top-k",
        "Solo cambia el tamaño del chunk",
    ], y=3.9, h=3.3, size=20)

    # 4. como se ve una ficha
    s = slide_titulo("Cómo se parte una ficha: ejemplo D05")
    datos = [["Condición", "Chunks de D05", "Passages que contiene cada chunk"]]
    for c in conds:
        d05 = [x for x in chunks[c] if x["doc_id"] == "D05"]
        datos.append([c, str(len(d05)), "; ".join(f"{x['chunk_id'].split('-')[1]}: P{pasajes(x)}".replace("PP", "P") for x in d05)])
    tabla(s, datos, 0.7, 1.6, 11.9, [2.0, 2.4, 7.5], size=16, alto=0.55)
    bullets(s, [
        "D05 es la ficha de los reactivos Mindray. Sus passages: P01 convocatoria, P02 ítem, P03 cronograma, P04 documentos, P05 postores, P06 buena pro, P07 contrato",
        "Con 100 palabras la buena pro (P06) queda en un chunk distinto al del ítem (P02)",
        "Con 180 y 320 palabras la ficha completa cabe en un solo chunk",
    ], y=4.2, h=2.8, size=19)
    notas(s, "Los chunks se arman con el chunker del Laboratorio 4. El id del chunk indica la ficha y el número de chunk dentro de ella.")

    # 5. tamaños reales
    s = slide_titulo("Resultados: los chunks que se generaron")
    ch = resumen["chunks"]
    datos = [["Condición", "Chunks", "Palabras (media)", "Palabras (máx.)", "Tokens (máx.)", "Cortados por tokens"]]
    for c in conds:
        datos.append([c, ch[c]["n_chunks"], round(ch[c]["palabras_media"]), ch[c]["palabras_max"],
                      ch[c]["tokens_max"], ch[c]["cortes_por_tokens"]])
    tabla(s, datos, 0.7, 1.6, 11.9, [2.0, 1.6, 2.3, 2.1, 1.9, 2.0], size=16, alto=0.55)
    tb = bullets(s, [
        "Interpretación",
        "  A mayor tamaño objetivo hay menos chunks, pero cada uno trae más texto",
        "  El texto del SEACE ocupa unos 2 tokens por palabra y el modelo e5 admite 512 tokens",
        f"  En large, {ch['large']['cortes_por_tokens']} de {ch['large']['n_chunks']} chunks se cerraron antes de 320 palabras; el más largo tiene {ch['large']['palabras_max']}",
        "  Así ningún chunk se trunca, pero large significa hasta 320 palabras o 512 tokens",
    ], y=4.3, h=3.0, size=20)
    tb.text_frame.paragraphs[0].runs[0].font.bold = True

    # 6. resultados recall
    s = slide_titulo("Resultados: Recall según el tamaño del chunk")
    cd = CategoryChartData()
    cd.categories = ["Recall@1", "Recall@3", "Recall@5"]
    for c in conds:
        cd.add_series(f"{c} ({palabras[c]})", [media(c, k) for k in (1, 3, 5)])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(1.5), Inches(7.4), Inches(5.0), cd)
    chart = gf.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    pl = chart.plots[0]
    pl.gap_width = 60
    for serie, c in zip(pl.series, conds):
        for punto, k in zip(serie.points, (1, 3, 5)):
            tf = punto.data_label.text_frame
            tf.text = num(media(c, k))
            tf.paragraphs[0].runs[0].font.size = Pt(11)
    chart.value_axis.maximum_scale = 0.8
    chart.value_axis.minimum_scale = 0
    chart.value_axis.has_major_gridlines = False
    datos = [["Condición", "R@1", "R@3", "R@5"]]
    for c in conds:
        datos.append([c] + [num(media(c, k)) for k in (1, 3, 5)])
    tabla(s, datos, 8.4, 1.7, 4.3, [1.5, 0.9, 0.9, 1.0], size=16)
    bullets(s, [
        "Métrica principal: Recall@3 medio",
        "baseline (180) es el mejor",
    ], x=8.4, y=3.8, w=4.4, h=3.0, size=18)

    # 7. interpretacion agregada
    s = slide_titulo("Interpretación: el resultado agregado")
    tb = bullets(s, [
        f"small (100) rinde claramente peor: Recall@3 de {num(media('small', 3))} contra {num(media('baseline', 3))} en baseline",
        "  Al partir las fichas en tantos chunks, la evidencia queda separada de su contexto",
        f"baseline (180) y large (320) están cerca ({num(media('baseline', 3))} y {num(media('large', 3))})",
        "  La diferencia es de un solo passage, no alcanza para decir que uno es mejor",
        "Más chunks no significa mejor recuperación",
        "  small tiene el doble de chunks que baseline y recupera menos",
        "El resultado vale para este corpus y estas consultas, no es una regla de chunking",
    ], size=21)

    # 8. por necesidad
    s = slide_titulo("Interpretación: por necesidad")
    nota = {"N1": "no cambia: encuentra la ficha Cisco, no la de Agilent",
            "N2": "0 en las tres: la frase \"o equivalente\" se diluye",
            "N3": "no cambia: encuentra la ficha Mindray, no la de Epson",
            "N4": "solo baseline trae el chunk con un solo postor",
            "N5": "el tamaño ayuda: con 100 palabras la buena pro queda aparte"}
    datos = [["Necesidad", "small", "baseline", "large", "Qué pasó"]]
    for i in ids:
        datos.append([i] + [num(r[(c, i)]["recall@3"]) for c in conds] + [nota[i]])
    tabla(s, datos, 0.7, 1.6, 11.9, [1.5, 1.2, 1.5, 1.2, 6.5], size=15, alto=0.5)
    bullets(s, [
        "Recall@3 por necesidad",
        "  3 de 5 necesidades no cambian con el tamaño; el efecto se concentra en N4 y N5",
    ], y=5.0, h=2.0, size=20)

    # 9. caso n5
    s = slide_titulo("Caso donde el tamaño ayuda: N5")
    t1 = {c: r[(c, "N5")]["top"][0] for c in conds}
    datos = [["Condición", "Primer resultado", "Passages que cubre", "Recall@3"]]
    for c in conds:
        t = t1[c]
        datos.append([c, f"{t['chunk_id']} ({t['score']:.3f})".replace(".", ","),
                      f"P{pasajes({'passage_ids': t['passage_ids']})}".replace("PP", "P"), num(r[(c, 'N5')]["recall@3"])])
    tabla(s, datos, 0.7, 1.6, 11.9, [2.0, 4.0, 3.9, 2.0], size=16, alto=0.55)
    bullets(s, [
        "N5: ¿Qué empresa ganó la compra de reactivos para el equipo Mindray del hospital de Jaén y por cuánto?",
        "  La evidencia es D05-P06, la buena pro",
        "Con 100 palabras el primer resultado es la ficha correcta, pero sin la buena pro",
        "  Documento correcto, fragmento incorrecto: la consulta describe el proceso, la evidencia es otra parte de la ficha",
        "Con 180 y 320 el chunk trae la ficha completa y el Recall sube a 1",
    ], y=4.0, h=3.3, size=19)

    # 10. caso n4
    s = slide_titulo("Caso donde el tamaño perjudica: N4")
    b4 = r[("baseline", "N4")]
    bullets(s, [
        "N4: ¿En qué procesos se presentó un solo postor?",
        f"  Recall@3: small {num(r[('small', 'N4')]['recall@3'])}, baseline {num(r[('baseline', 'N4')]['recall@3'])}, large {num(r[('large', 'N4')]['recall@3'])}",
        "En baseline el chunk D07-C02 (postores, buena pro y contrato de Vaso de Leche) entra en el puesto 3",
        "En large ese passage queda dentro del chunk de la ficha completa",
        "  El chunk baja al puesto 11 y sale del top-3",
        "  Posible explicación: en la ficha completa pesan más la convocatoria y los alimentos que la frase un solo postor",
        "Un chunk más grande no siempre ayuda: puede diluir una evidencia corta",
    ], size=20)

    # 11. conclusion
    s = slide_titulo("Conclusión y limitación")
    bullets(s, [
        "Con este corpus, estas consultas y estos qrels, el tamaño del chunk cambió el Recall@3 medio de "
        f"{num(media('small', 3))} (100) a {num(media('baseline', 3))} (180) y {num(media('large', 3))} (320)",
        "  La diferencia principal está entre small y los otros dos, y la explican N5 y N4",
        "El tamaño ayuda cuando la evidencia y lo que la identifica deben quedar juntos, y perjudica cuando la evidencia es corta y se diluye",
        "Limitaciones",
        "  Solo 5 necesidades: un passage mueve el Recall medio entre 0,05 y 0,1",
        "  large está limitado por los 512 tokens del modelo",
        "  No se puede afirmar que un tamaño sea el mejor en general",
    ], size=21)


def construir():
    p = nueva_presentacion()
    agregar(p)
    salida = aqui / "Experimento2.pptx"
    p.save(salida)
    exportar(salida, aqui / "Experimento2.pdf", aqui / "_png_exp2")
    print("listo", salida)


if __name__ == "__main__":
    construir()
